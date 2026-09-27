"""
eCourts Uttarakhand - Data Crawler (FINAL)
==========================================
Run:  python crawler.py
"""

import sys, io, time
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from pathlib import Path
from datetime import datetime

try:
    from selenium import webdriver
    from selenium.webdriver.chrome.options import Options
    from selenium.webdriver.chrome.service import Service
    from selenium.webdriver.common.by import By
    from selenium.common.exceptions import (
        NoSuchElementException, InvalidSessionIdException,
        WebDriverException
    )
    from webdriver_manager.chrome import ChromeDriverManager
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    print("ERROR: Run this first:")
    print("  pip install selenium webdriver-manager openpyxl")
    sys.exit(1)

OUTPUT = "ecourts_dataset.xlsx"

SECTIONS = [
    ("Case Type Search",
     "https://hcservices.ecourts.gov.in/ecourtindiaHC/cases/s_casetype.php?state_cd=15&dist_cd=1&court_code=1&stateNm=Uttarakhand"),
    ("Party Name Search",
     "https://hcservices.ecourts.gov.in/ecourtindiaHC/cases/ki_petres.php?state_cd=15&dist_cd=1&court_code=1&stateNm=Uttarakhand"),
    ("Orders by Date",
     "https://hcservices.ecourts.gov.in/ecourtindiaHC/cases/s_orderdate.php?state_cd=15&dist_cd=1&court_code=1&stateNm=Uttarakhand"),
    ("Orders by Judge",
     "https://hcservices.ecourts.gov.in/ecourtindiaHC/cases/s_order.php?state_cd=15&dist_cd=1&court_code=1&stateNm=Uttarakhand"),
]

# ---------------------------------------------------------------------------
def start_browser():
    opts = Options()
    opts.add_argument("--disable-blink-features=AutomationControlled")
    opts.add_experimental_option("excludeSwitches", ["enable-automation"])
    opts.add_experimental_option("useAutomationExtension", False)
    opts.add_argument(
        "user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    )
    svc    = Service(ChromeDriverManager().install())
    driver = webdriver.Chrome(service=svc, options=opts)
    driver.maximize_window()
    return driver


def safe_get(driver, url):
    """Navigate to URL, restart browser if session died."""
    try:
        driver.get(url)
        return driver
    except (InvalidSessionIdException, WebDriverException):
        print("  Browser was closed. Restarting Chrome...")
        try:
            driver.quit()
        except:
            pass
        driver = start_browser()
        driver.get(url)
        return driver


# ---------------------------------------------------------------------------
def get_table_data(driver):
    time.sleep(2)
    all_rows = []

    try:
        tables = driver.find_elements(By.TAG_NAME, "table")
    except (InvalidSessionIdException, WebDriverException):
        return []

    tables = [t for t in tables if len(t.find_elements(By.TAG_NAME, "tr")) > 2]

    for t_idx, table in enumerate(tables):
        try:
            rows = table.find_elements(By.TAG_NAME, "tr")
            if not rows:
                continue

            # Find headers
            headers, data_start = [], 1
            for i, row in enumerate(rows[:3]):
                ths = row.find_elements(By.TAG_NAME, "th")
                if ths:
                    headers    = [h.text.strip().replace("\n"," ") or f"Col{j+1}" for j,h in enumerate(ths)]
                    data_start = i + 1
                    break
            if not headers:
                tds0    = rows[0].find_elements(By.TAG_NAME, "td")
                headers = [c.text.strip() or f"Col{i+1}" for i,c in enumerate(tds0)]
                if not any(headers):
                    continue

            # Deduplicate header names
            seen, clean = {}, []
            for h in headers:
                n = seen.get(h, 0); seen[h] = n+1
                clean.append(h if n == 0 else f"{h}_{n}")

            # Data rows
            for row in rows[data_start:]:
                tds   = row.find_elements(By.TAG_NAME, "td")
                cells = [td.text.strip() for td in tds]
                if not any(cells):
                    continue
                while len(cells) < len(clean): cells.append("")
                cells = cells[:len(clean)]
                rec = dict(zip(clean, cells))
                rec["_table"] = t_idx
                all_rows.append(rec)
        except Exception:
            continue

    return all_rows


def find_next(driver):
    xpaths = [
        "//a[normalize-space(text())='Next']",
        "//a[normalize-space(text())='next']",
        "//a[normalize-space(text())='>']",
        "//a[normalize-space(text())='>>']",
        "//input[@value='Next']",
        "//button[contains(translate(text(),'NEXT','next'),'next')]",
        "//li[contains(@class,'next') and not(contains(@class,'disabled'))]/a",
        "//a[@aria-label='Next']",
        "//a[contains(@class,'next') and not(contains(@class,'disabled'))]",
        "//span[@class='next']/a",
    ]
    for xp in xpaths:
        try:
            el = driver.find_element(By.XPATH, xp)
            if el.is_displayed() and el.is_enabled():
                return el
        except NoSuchElementException:
            pass
    return None


def collect_pages(driver, label):
    data, page = [], 1
    while True:
        print(f"  Page {page}...", end=" ", flush=True)
        rows = get_table_data(driver)
        if rows:
            for r in rows:
                r["_search"] = label
                r["_page"]   = page
                r["_time"]   = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            data.extend(rows)
            print(f"{len(rows)} rows  (total: {len(data)})")
        else:
            print("0 rows.")
            ans = input("  See data in Chrome but got 0? Type 'c'+Enter to retry, else Enter to stop: ").strip().lower()
            if ans != "c":
                break

        btn = find_next(driver)
        if btn:
            try:
                driver.execute_script("arguments[0].scrollIntoView(true);", btn)
                time.sleep(0.3)
                driver.execute_script("arguments[0].click();", btn)
                time.sleep(2)
                page += 1
            except Exception as e:
                print(f"  Next page click failed: {e}")
                break
        else:
            print(f"  All {page} page(s) collected.")
            break
    return data


# ---------------------------------------------------------------------------
def save_excel(rows, filename=OUTPUT):
    wb = Workbook()
    ws = wb.active
    ws.title = "Case Data"

    if not rows:
        ws.append(["No data collected."])
        wb.save(filename)
        return 0

    keys    = list(dict.fromkeys(k for r in rows for k in r.keys()))
    visible = [k for k in keys if not k.startswith("_")]
    hidden  = [k for k in keys if k.startswith("_")]
    cols    = visible + hidden

    ws.append(cols)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.font = Font(color="FFFFFF", bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28

    for i, rec in enumerate(rows, start=2):
        ws.append([rec.get(c,"") for c in cols])
        fc = "DCE6F1" if i%2==0 else "FFFFFF"
        for cell in ws[i]:
            cell.fill = PatternFill("solid", fgColor=fc)
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        w = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(w+4, 55)
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Summary")
    ws2.append(["Info","Value"])
    ws2.append(["Total rows", len(rows)])
    ws2.append(["Created at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws2.append(["Source", "https://hcservices.ecourts.gov.in/ecourtindiaHC/"])
    for cell in ws2[1]:
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.font = Font(color="FFFFFF", bold=True)
    ws2.column_dimensions["A"].width = 20
    ws2.column_dimensions["B"].width = 50

    wb.save(filename)
    return len(rows)


# ---------------------------------------------------------------------------
def main():
    print("=" * 60)
    print("  eCourts Uttarakhand - Data Crawler")
    print("=" * 60)
    print()
    print("  HOW TO USE:")
    print("  -----------")
    print("  Chrome will open for each search section.")
    print("  YOU fill the form + solve CAPTCHA + click Search.")
    print("  Script collects ALL pages of results automatically.")
    print("  Excel file saved at the end.")
    print()
    print("  Press Enter to open Chrome now...")
    input()

    driver    = start_browser()
    collected = []

    try:
        for num, (name, url) in enumerate(SECTIONS, 1):
            print(f"\n{'='*60}")
            print(f"  [{num}/{len(SECTIONS)}] {name}")
            print(f"{'='*60}")
            print("  Press 's'+Enter to SKIP, or just Enter to open this page:")
            if input("  > ").strip().lower() == "s":
                print("  Skipped.")
                continue

            driver = safe_get(driver, url)
            time.sleep(3)

            while True:
                print()
                print("  >>>  DO THIS IN CHROME:  <<<")
                print("  1. Fill the search form")
                print("  2. Solve the CAPTCHA (type letters shown)")
                print("  3. Click Search / Submit / Go")
                print("  4. Wait for results table to appear")
                print()
                input("  Press Enter here AFTER results load in Chrome > ")

                lbl  = input(f"  Label for this data (e.g. 'Writ Petition 2024') [{name}]: ").strip() or name
                rows = collect_pages(driver, lbl)
                collected.extend(rows)
                print(f"  Got {len(rows)} rows. Grand total: {len(collected)}")

                print()
                print("  Search AGAIN with different filter (another case type / date)?")
                if input("  Type 'y'+Enter for yes, Enter to move on: ").strip().lower() != "y":
                    break
                driver = safe_get(driver, url)
                time.sleep(3)

    except KeyboardInterrupt:
        print("\n\n  Stopped - saving data...")
    finally:
        try:
            driver.quit()
        except:
            pass
        print("  Browser closed.")

    print(f"\n  Saving to {OUTPUT}...")
    total = save_excel(collected)
    path  = Path(OUTPUT).resolve()
    print()
    print("=" * 60)
    print("  DONE!")
    print(f"  File  : {path}")
    print(f"  Rows  : {total}")
    print("=" * 60)


if __name__ == "__main__":
    main()
